# Магазин «Сулайман-Тоо» — разовый импорт истории долга с поставщиком (Икром) из Excel.
# Чистые функции: разбор листа и метки для защиты от повторного импорта.
from datetime import date, datetime
import re

SEED_PATH = "data/ikrom_ledger_seed.xlsx"
IMPORT_NAME = "импорт Excel"
_MARK_RE = re.compile(r"\[excel:[^\]]+\]")


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def marker(d: date, typ: str, amount: float) -> str:
    return f"[excel:{d.isoformat()}:{typ}:{round(float(amount), 2):g}]"


def created_at_for(d: date) -> str:
    """Полдень по Бишкеку (UTC+6) — чтобы история шла по настоящим датам."""
    return f"{d.isoformat()}T12:00:00+06:00"


def parse_rows(rows):
    """rows: итерируемое кортежей (A..G). Возвращает список записей для supplier_ledger."""
    out = []
    for r in rows:
        r = list(r) + [None] * 7
        d = _to_date(r[0])
        if d is None:
            continue  # строки «итого за месяц» и заголовки
        total = _num(r[3])
        if total > 0:
            out.append({"type": "товар", "amount": round(total, 2), "date": d, "note": "ТМЗ за день"})
        pay = _num(r[4])
        note = " ".join(str(x).strip() for x in (r[5], r[6]) if x not in (None, "")).strip()
        if pay > 0 and "икром" in note.lower():
            typ = "возврат" if "возврат" in note.lower() else "оплата"
            out.append({"type": typ, "amount": round(pay, 2), "date": d, "note": note})
    for e in out:
        e["marker"] = marker(e["date"], e["type"], e["amount"])
        e["comment"] = f"{e['marker']} {e['date'].strftime('%d.%m.%Y')} · {e['note']}"
        e["side"] = "поставщик" if e["type"] == "товар" else "магазин"
    return out


def parse_workbook(src, sheet="Лист1"):
    from openpyxl import load_workbook
    wb = load_workbook(src, data_only=True, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    return parse_rows(ws.iter_rows(values_only=True))


def existing_markers(entries) -> set:
    s = set()
    for e in entries or []:
        s.update(_MARK_RE.findall(e.get("comment") or ""))
    return s


def new_only(parsed, entries):
    have = existing_markers(entries)
    return [e for e in parsed if e["marker"] not in have]


def summary(parsed) -> dict:
    res = {t: {"count": 0, "sum": 0.0} for t in ("товар", "оплата", "возврат")}
    for e in parsed:
        res[e["type"]]["count"] += 1
        res[e["type"]]["sum"] += e["amount"]
    return res


def to_rows(parsed, accepted=True, user=None, now_iso=None):
    user = user or {}
    return [{
        "type": e["type"], "amount": float(e["amount"]), "comment": e["comment"],
        "created_at": created_at_for(e["date"]),
        "created_by_user_id": user.get("id") or None, "created_by_name": IMPORT_NAME,
        "created_by_role": user.get("role"), "created_by_side": e["side"],
        "accepted": bool(accepted),
        "accepted_at": now_iso if accepted else None,
        "accepted_by_user_id": (user.get("id") or None) if accepted else None,
        "accepted_by_name": (user.get("username") or IMPORT_NAME) if accepted else None,
    } for e in parsed]
