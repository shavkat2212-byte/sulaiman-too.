# Тесты бланка обязательства: python -m pytest tests/
import io, os, sys
from datetime import date
import pytest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from obligation import (build_docx, build_obligation_data, number_words, render_html, schedule_layout,  # noqa: E402
                        seller_settings, som_words)


@pytest.mark.parametrize("n,text", [
    (0, "ноль сомов"), (1, "один сом"), (2, "два сома"), (4, "четыре сома"), (5, "пять сомов"),
    (11, "одиннадцать сомов"), (12, "двенадцать сомов"), (21, "двадцать один сом"), (22, "двадцать два сома"),
    (100, "сто сомов"), (111, "сто одиннадцать сомов"), (1000, "одна тысяча сомов"), (2000, "две тысячи сомов"),
    (5000, "пять тысяч сомов"), (11000, "одиннадцать тысяч сомов"), (21000, "двадцать одна тысяча сомов"),
    (55250, "пятьдесят пять тысяч двести пятьдесят сомов"), (1000000, "один миллион сомов"),
    (2345678, "два миллиона триста сорок пять тысяч шестьсот семьдесят восемь сомов"),
])
def test_som_words(n, text):
    assert som_words(n) == text


def test_tyiyn_and_negative():
    assert som_words(12.5) == "двенадцать сомов 50 тыйынов"
    assert number_words(-3) == "минус три"


SALE = {"id": "1610261200", "day": "2026-10-16", "client_id": 1, "down_payment": 5000, "credit_balance": 33040,
        "total_sale": 30000, "name": "Договор рассрочки №1610 [Холодильник lg (1 шт.), Чайник (2 шт.)] — Иванов Иван"}
CLIENT = {"id": 1, "fio": "Иванов Иван", "phone": "0555 12-34-56", "address": "г. Ош, ул. Ленина 1"}


def pays(n, amount=5506):
    return [{"id": i, "sale_id": SALE["id"], "due_date": f"{2026 + (10 + i) // 12}-{(10 + i) % 12 + 1:02d}-16",
             "amount_expected": amount} for i in range(n)]


def test_data_fill():
    d = build_obligation_data(SALE, CLIENT, pays(6), {"inn": "123", "address": "", "city": "Ош"})
    assert d["number"] == "1610261200" and d["date"] == date(2026, 10, 16)
    assert d["goods_name"] == "Холодильник lg, Чайник" and d["goods_qty"] == 3
    assert d["total"] == 38040 and d["down"] == 5000 and d["remaining"] == 33040   # цена покупки (total_sale) не показывается
    assert d["remaining_words"] == "тридцать три тысячи сорок сомов"
    assert len(d["schedule"]) == 6 and d["schedule_total"] == 6 * 5506
    assert d["schedule"][0]["date"] == date(2026, 11, 16)
    assert d["pin"] == "" and d["passport_date"] is None   # колонок нет — пустые линии


def test_schedule_layout():
    assert schedule_layout(0) == 6 and schedule_layout(12) == 6 and schedule_layout(13) == 7 and schedule_layout(24) == 12


def test_html_has_values_and_print():
    d = build_obligation_data(SALE, {**CLIENT, "pin": "12345678901234", "passport_date": "2020-03-15"}, pays(14), {})
    h = render_html(d)
    for s in ("window.print()", "@media print", "Иванов Иван", "12345678901234", "«15» марта 2020 г.",
              "33 040", "<td class=\"n\">14</td>", "Итого:", "ОБЯЗАТЕЛЬСТВО ПО ОПЛАТЕ ТОВАРА"):
        assert s in h, s
    assert "30 000" not in h   # закупочная цена товара не попадает в бланк


def test_html_escapes():
    d = build_obligation_data(SALE, {**CLIENT, "fio": "<script>x</script>"}, [], {})
    assert "<script>x" not in render_html(d)


def test_docx_builds():
    from docx import Document
    d = build_obligation_data(SALE, CLIENT, pays(14), {"inn": "", "address": "", "city": ""})
    doc = Document(io.BytesIO(build_docx(d)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ОБЯЗАТЕЛЬСТВО ПО ОПЛАТЕ ТОВАРА" in text and "Иванов Иван" in text
    sch = doc.tables[3]
    assert len(sch.rows) == 7 + 2   # 14 платежей → 7+7 строк + шапка + итого
    assert sch.cell(1, 2).text == "5 506"


def test_seller_settings_optional():
    assert seller_settings({}) == {"inn": "20803200101311947", "address": "г. Ош, ул. Амир Темур 112", "city": "Ош"}
    assert seller_settings({"SELLER_INN": "", "SELLER_ADDRESS": "ул. Другая 1"})["inn"] == "20803200101311947"
    assert seller_settings({"SELLER_INN": "0101", "SELLER_CITY": "Ош"})["city"] == "Ош"
