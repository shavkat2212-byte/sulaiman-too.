# Бланк «Обязательство по оплате товара (в рассрочку)»: HTML для печати + Word. PDF не делаем.
import html
import io
import math
from datetime import date, datetime

SELLER_NAME = "ИП «Ташходжаев А.А.»"
SELLER_DEFAULTS = {"inn": "20803200101311947", "address": "г. Ош, ул. Амир Темур 112", "city": "Ош"}  # st.secrets может переопределить
MONTHS_GEN = ["", "января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
              "сентября", "октября", "ноября", "декабря"]

# ---------------- сумма прописью ----------------
_ONES_M = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
_ONES_F = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
_TEENS = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", "пятнадцать",
          "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
_TENS = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
_HUNDREDS = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот"]
# (формы: 1, 2-4, 5+), женский род?
_SCALES = [(("", "", ""), False), (("тысяча", "тысячи", "тысяч"), True),
           (("миллион", "миллиона", "миллионов"), False), (("миллиард", "миллиарда", "миллиардов"), False)]


def plural(n, forms):
    n = abs(int(n)) % 100
    if 11 <= n <= 19:
        return forms[2]
    n %= 10
    if n == 1:
        return forms[0]
    if 2 <= n <= 4:
        return forms[1]
    return forms[2]


def _triad(n, feminine):
    words = [_HUNDREDS[n // 100]]
    t = n % 100
    if 10 <= t <= 19:
        words.append(_TEENS[t - 10])
    else:
        words.append(_TENS[t // 10])
        words.append((_ONES_F if feminine else _ONES_M)[t % 10])
    return [w for w in words if w]


def number_words(n):
    """Целое число прописью, мужской род: 1 → «один», 21000 → «двадцать одна тысяча»."""
    n = int(n)
    if n == 0:
        return "ноль"
    if n < 0:
        return "минус " + number_words(-n)
    parts, i = [], 0
    while n > 0:
        if i >= len(_SCALES):
            raise ValueError("Слишком большое число")
        tri = n % 1000
        if tri:
            forms, fem = _SCALES[i]
            w = _triad(tri, fem)
            if i:
                w.append(plural(tri, forms))
            parts = w + parts
        n //= 1000
        i += 1
    return " ".join(parts)


def som_words(amount):
    """Сумма прописью: 55250 → «пятьдесят пять тысяч двести пятьдесят сомов»; тыйыны — цифрами."""
    amount = round(float(amount or 0), 2)
    whole = int(amount)
    tyiyn = int(round((amount - whole) * 100))
    text = f"{number_words(whole)} {plural(whole, ('сом', 'сома', 'сомов'))}"
    if tyiyn:
        text += f" {tyiyn:02d} {plural(tyiyn, ('тыйын', 'тыйына', 'тыйынов'))}"
    return text


# ---------------- данные ----------------
def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _day(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt, n in (("%Y-%m-%d", 10), ("%d.%m.%Y", 10)):
        try:
            return datetime.strptime(s[:n], fmt).date()
        except ValueError:
            pass
    return None


def money(v):
    v = _num(v)
    txt = f"{v:,.2f}".rstrip("0").rstrip(".") if v != int(v) else f"{int(v):,}"
    return txt.replace(",", " ")


def seller_settings(secrets=None):
    out = dict(SELLER_DEFAULTS)
    try:
        if secrets is None:
            import streamlit as st
            secrets = st.secrets
        for key, name in (("inn", "SELLER_INN"), ("address", "SELLER_ADDRESS"), ("city", "SELLER_CITY")):
            val = str(secrets.get(name, "") or "").strip()
            if val:
                out[key] = val
    except Exception:
        pass
    return out


def build_obligation_data(sale, client, payments, seller=None):
    from installment_logic import parse_goods
    client = client or {}
    seller = seller if seller is not None else dict(SELLER_DEFAULTS)
    goods = parse_goods(sale.get("name"))
    d = _day(sale.get("day")) or _day(sale.get("date"))
    down = _num(sale.get("down_payment"))
    credit = _num(sale.get("credit_balance"))
    rows = sorted([p for p in payments if str(p.get("sale_id")) == str(sale.get("id"))],
                  key=lambda p: (_day(p.get("due_date")) or date.max, str(p.get("id"))))
    schedule = [{"date": _day(p.get("due_date")), "amount": _num(p.get("amount_expected"))} for p in rows]
    pdate = _day(client.get("passport_date"))
    return {
        "number": str(sale.get("id") or ""),
        "date": d,
        "city": seller.get("city", ""),
        "seller_inn": seller.get("inn", ""),
        "seller_address": seller.get("address", ""),
        "fio": str(client.get("fio") or "").strip(),
        "phone": str(client.get("phone") or "").strip(),
        "address": str(client.get("address") or "").strip(),
        "pin": str(client.get("pin") or "").strip(),
        "passport": str(client.get("passport") or "").strip(),
        "passport_issued_by": str(client.get("passport_issued_by") or "").strip(),
        "passport_date": pdate,
        "goods_name": ", ".join(n for n, _ in goods),
        "goods_qty": sum(q for _, q in goods) if goods else None,
        "total": down + credit,          # цена для клиента = взнос + сумма рассрочки
        "down": down,
        "remaining": credit,
        "remaining_words": som_words(credit),
        "schedule": schedule,
        "schedule_total": sum(r["amount"] for r in schedule),
    }


def schedule_layout(count):
    """Сколько строк в каждой из двух половин таблицы: минимум 6+6 как в бланке."""
    return max(6, math.ceil(count / 2))


def date_parts(d):
    if not d:
        return None
    return f"«{d.day:02d}» {MONTHS_GEN[d.month]} {d.year} г."


# ---------------- HTML ----------------
def _f(value, width_em=10):
    """Заполненное поле — подчёркнутое значение; пустое — линия для ручки."""
    if value in (None, ""):
        return f'<span class="blank" style="min-width:{width_em}em"></span>'
    return f'<span class="fill">{html.escape(str(value))}</span>'


CSS = """
*{box-sizing:border-box}
body{margin:0;background:#e9e9e9;font-family:'Times New Roman',Times,serif;font-size:10.5pt;color:#000}
.toolbar{position:sticky;top:0;background:#fff;padding:8px;text-align:center;border-bottom:1px solid #ccc;z-index:2}
.toolbar button{font-size:16px;padding:8px 22px;border-radius:8px;border:0;background:#ff4b4b;color:#fff;cursor:pointer}
.sheet{width:210mm;min-height:297mm;margin:10px auto;background:#fff;padding:13mm 15mm;box-shadow:0 0 6px rgba(0,0,0,.25)}
.wrap{width:100%;overflow-x:auto}
p{margin:0 0 3pt 0;line-height:1.25}
.row{display:flex;justify-content:space-between;gap:8mm}.row>div{flex:1}
.right{text-align:right}
h1{font-size:16pt;text-align:center;margin:8pt 0 0 0}
h2{font-size:13pt;text-align:center;margin:0 0 6pt 0}
h3{font-size:10.5pt;margin:8pt 0 4pt 0}
.fill{border-bottom:1px solid #000;padding:0 4px;font-style:italic}
.blank{display:inline-block;border-bottom:1px solid #000;height:1em;vertical-align:baseline}
.grow{display:flex;align-items:baseline;gap:4px}.grow .blank,.grow .fill{flex:1}
table.sch{border-collapse:collapse;width:100%;table-layout:fixed}
table.sch td,table.sch th{border:1px solid #000;padding:1px 3px;font-size:10pt;height:%(rowh)s}
table.sch th{font-size:9pt;height:auto}
table.sch td.n{text-align:center}table.sch td.a{text-align:right}
table.sch td.gap{border:0;width:0.5cm}
table.sch td.none{border:0}
.just{text-align:justify}
.small{font-size:8pt}.note{font-size:9.5pt;font-style:italic;margin-top:6pt}
.sig{margin-top:10pt}
@page{size:A4;margin:0}
@media print{
  body{background:#fff}
  .toolbar{display:none}
  .sheet{margin:0;box-shadow:none;width:210mm;min-height:auto}
  .wrap{overflow:visible}
}
"""


def render_html(d, toolbar=True):
    sch = d["schedule"]
    half = schedule_layout(len(sch))
    rowh = "0.75cm" if half <= 6 else ("0.6cm" if half <= 9 else "0.5cm")
    css = CSS.replace("%(rowh)s", rowh)

    def cells(i):
        if i < len(sch):
            r = sch[i]
            return (f'<td class="n">{i + 1}</td><td class="n">{r["date"].strftime("%d.%m.%Y") if r["date"] else ""}</td>'
                    f'<td class="a">{money(r["amount"])}</td><td></td>')
        return f'<td class="n">{i + 1}</td><td></td><td></td><td></td>'

    hdr = "<th style='width:0.7cm'>№</th><th style='width:2.2cm'>Дата платежа</th><th style='width:2cm'>Сумма, сом</th><th>Отметка об оплате (подпись продавца)</th>"
    body = "".join(f"<tr>{cells(i)}<td class='gap'></td>{cells(i + half)}</tr>" for i in range(half))
    body += (f"<tr><td class='none' colspan='5'></td><td colspan='2' class='a'><b>Итого:</b></td>"
             f"<td class='a'><b>{money(d['schedule_total']) if sch else ''}</b></td><td></td></tr>")
    passport_line = (f'<p class="grow">Паспорт / ID: {_f(d["passport"])}</p>' if d["passport"]
                     else '<p>Паспорт / ID: серия <span class="blank" style="min-width:4em"></span> № <span class="blank" style="min-width:8em"></span></p>')
    date_txt = date_parts(d["date"])
    btn = ('<div class="toolbar"><button onclick="window.print()">🖨 Печать</button></div>' if toolbar else "")
    return f"""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Обязательство № {html.escape(d['number'])}</title>
<style>{css}</style></head><body>{btn}<div class="wrap"><div class="sheet">
<div class="row"><div>№ {_f(d['number'], 6)} &nbsp;&nbsp; г. {_f(d['city'], 9)}</div>
<div class="right">Дата: {_f(date_txt) if date_txt else '«____» ______________ 20___ г.'}</div></div>
<h1>ОБЯЗАТЕЛЬСТВО ПО ОПЛАТЕ ТОВАРА</h1><h2>(В РАССРОЧКУ)</h2>
<p><b>Продавец:</b> {SELLER_NAME}, ИНН {_f(d['seller_inn'], 12)}</p>
<p class="grow">Адрес продавца: {_f(d['seller_address'])}</p>
<h3>1. Данные покупателя</h3>
<p class="grow">Ф. И. О.: {_f(d['fio'])}</p>
<div class="row"><div><p class="grow">ПИН (14 цифр): {_f(d['pin'])}</p></div><div><p class="grow">Телефон: {_f(d['phone'])}</p></div></div>
<div class="row"><div>{passport_line}</div><div><p>Дата выдачи: {_f(date_parts(d['passport_date'])) if d['passport_date'] else '«____» ____________ 20___ г.'}</p></div></div>
<p class="grow">Кем выдан: {_f(d['passport_issued_by'])}</p>
<p class="grow">Адрес проживания: {_f(d['address'])}</p>
<h3>2. Данные о товаре</h3>
<p class="grow">Наименование / модель: {_f(d['goods_name'])} <span>Количество: {_f(d['goods_qty'], 3)} шт.</span></p>
<div class="row"><div><p>Общая стоимость товара: {_f(money(d['total']), 7)} сом</p></div><div><p>Первоначальный взнос: {_f(money(d['down']), 7)} сом</p></div></div>
<p><b>Остаток к оплате (сумма рассрочки):</b> {_f(money(d['remaining']), 7)} сом (прописью: {_f(d['remaining_words'])}).</p>
<h3>3. График платежей</h3>
<table class="sch"><colgroup><col style="width:0.7cm"><col style="width:2.2cm"><col style="width:2cm"><col>
<col style="width:0.5cm"><col style="width:0.7cm"><col style="width:2.2cm"><col style="width:2cm"><col></colgroup>
<tr>{hdr}<td class="gap"></td>{hdr}</tr>{body}</table>
<h3>4. Обязательства сторон</h3>
<p class="just">4.1. Покупатель подтверждает получение указанного товара в надлежащем состоянии, согласен с его стоимостью и обязуется вносить платежи в сроки и суммах согласно графику.</p>
<p class="just">4.2. Если Покупатель пропустит два платежа подряд, Продавец вправе потребовать погашения всего оставшегося долга сразу.</p>
<p class="just">4.3. При неоплате споры разрешаются в суде по месту нахождения Продавца в соответствии с законодательством Кыргызской Республики.</p>
<p class="just">4.4. Покупатель обязуется сообщить Продавцу о смене телефона или адреса. Уведомления, направленные по указанному телефону (звонок, SMS, WhatsApp), считаются полученными.</p>
<h3>5. Подтверждение и подписи</h3>
<p>С условиями ознакомлен(а), товар получил(а), подписываю добровольно.</p>
<div class="row sig"><div><b>Покупатель:</b></div><div><b>Продавец:</b> {SELLER_NAME}</div></div>
<div class="row sig"><div>_____________ /{_f(d['fio'], 11) if d['fio'] else '______________________'}/</div><div>_____________ /______________________/</div></div>
<div class="row small"><div>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;подпись&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;Ф. И. О.</div><div>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;подпись&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;Ф. И. О.</div></div>
<p class="sig">Дата подписания: «____» ______________ 20___ г.</p>
<p class="note">Составлено в двух экземплярах, по одному для каждой стороны.</p>
</div></div></body></html>"""


# ---------------- Word ----------------
def build_docx(d):
    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    F, SZ = "Times New Roman", 10.5
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.5)
    s.top_margin = s.bottom_margin = Cm(1.3)
    stl = doc.styles["Normal"]
    stl.font.name, stl.font.size = F, Pt(SZ)
    stl.element.rPr.rFonts.set(qn("w:eastAsia"), F)
    pf = stl.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0

    def P(parts, align=None, after=0, before=0, size=None, cell=None):
        if cell is not None:
            ps = cell.paragraphs
            p = ps[0] if (len(ps) == 1 and not ps[0].runs) else cell.add_paragraph()
        else:
            p = doc.add_paragraph()
        if isinstance(parts, str):
            parts = [(parts, False)]
        for part in parts:
            t, b = part[0], part[1]
            r = p.add_run(t)
            r.bold = b
            r.font.name, r.font.size = F, Pt(size or SZ)
            if len(part) > 2 and part[2]:  # заполненное поле
                r.underline, r.italic = True, True
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
        return p

    def V(value, blank="______________"):
        return (f" {value} ", False, True) if value not in (None, "") else (blank, False)

    def H(t):
        return P([(t, True)], before=8, after=4)

    def borders(tbl, on=True):
        b = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            x = OxmlElement("w:" + e)
            x.set(qn("w:val"), "single" if on else "nil")
            if on:
                x.set(qn("w:sz"), "4")
                x.set(qn("w:color"), "000000")
            b.append(x)
        tbl._tbl.tblPr.append(b)

    def widths(tbl, ws):
        tbl.autofit = False
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), ws):
            gc.set(qn("w:w"), str(int(w * 567)))
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tbl._tbl.tblPr.append(lay)
        for row in tbl.rows:
            for c, w in zip(row.cells, ws):
                c.width = Cm(w)

    def cell_borders(c, edges):
        tcPr = c._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for e in edges:
            x = OxmlElement("w:" + e)
            x.set(qn("w:val"), "nil")
            b.append(x)
        tcPr.append(b)

    date_txt = date_parts(d["date"])
    t = doc.add_table(rows=1, cols=2); borders(t, False); widths(t, [9, 9])
    P([("№ ", False), V(d["number"], "________"), ("          г. ", False), V(d["city"], "________________")], cell=t.cell(0, 0))
    P([("Дата: ", False), V(date_txt, "«____» ______________ 20___ г.")], align=A.RIGHT, cell=t.cell(0, 1))
    P([("ОБЯЗАТЕЛЬСТВО ПО ОПЛАТЕ ТОВАРА", True)], align=A.CENTER, before=8, size=16)
    P([("(В РАССРОЧКУ)", True)], align=A.CENTER, after=6, size=13)
    P([("Продавец: ", True), (f"{SELLER_NAME}, ИНН ", False), V(d["seller_inn"], "______________________")], after=4)
    P([("Адрес продавца: ", False), V(d["seller_address"], "_" * 76)])
    H("1. Данные покупателя")
    P([("Ф. И. О.: ", False), V(d["fio"], "_" * 83)], after=4)
    t = doc.add_table(rows=2, cols=2); borders(t, False); widths(t, [9, 9])
    P([("ПИН (14 цифр): ", False), V(d["pin"], "_" * 30)], cell=t.cell(0, 0), after=4)
    P([("Телефон: ", False), V(d["phone"], "_" * 35)], cell=t.cell(0, 1), after=4)
    if d["passport"]:
        P([("Паспорт / ID: ", False), V(d["passport"])], cell=t.cell(1, 0), after=4)
    else:
        P("Паспорт / ID: серия ______ № _______________", cell=t.cell(1, 0), after=4)
    P([("Дата выдачи: ", False), V(date_parts(d["passport_date"]), "«____» ____________ 20___ г.")], cell=t.cell(1, 1), after=4)
    P([("Кем выдан: ", False), V(d["passport_issued_by"], "_" * 81)], after=4)
    P([("Адрес проживания: ", False), V(d["address"], "_" * 74)])
    H("2. Данные о товаре")
    P([("Наименование / модель: ", False), V(d["goods_name"], "_" * 47), (" Количество: ", False),
       V(d["goods_qty"], "______"), (" шт.", False)], after=4)
    t = doc.add_table(rows=1, cols=2); borders(t, False); widths(t, [9, 9])
    P([("Общая стоимость товара: ", False), V(money(d["total"])), (" сом", False)], cell=t.cell(0, 0))
    P([("Первоначальный взнос: ", False), V(money(d["down"])), (" сом", False)], cell=t.cell(0, 1))
    P([("Остаток к оплате (сумма рассрочки): ", True), V(money(d["remaining"])), (" сом (прописью: ", False),
       V(d["remaining_words"]), (").", False)], before=3, after=4)

    H("3. График платежей")
    sch, n = d["schedule"], schedule_layout(len(d["schedule"]))
    t = doc.add_table(rows=n + 2, cols=9); borders(t, True)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths(t, [0.7, 2.2, 2.0, 3.6, 0.5, 0.7, 2.2, 2.0, 3.6])
    for j, h in enumerate(["№", "Дата платежа", "Сумма, сом", "Отметка об оплате (подпись продавца)"]):
        for off in (0, 5):
            P([(h, True)], align=A.CENTER, cell=t.cell(0, j + off), size=9)
    for i in range(n):
        for off, k in ((0, i), (5, i + n)):
            P(str(k + 1), align=A.CENTER, cell=t.cell(i + 1, off))
            if k < len(sch):
                r = sch[k]
                P(r["date"].strftime("%d.%m.%Y") if r["date"] else "", align=A.CENTER, cell=t.cell(i + 1, off + 1))
                P(money(r["amount"]), align=A.RIGHT, cell=t.cell(i + 1, off + 2))
    rh = Cm(0.75 if n <= 6 else (0.6 if n <= 9 else 0.5))
    for r in t.rows[1:]:
        r.height, r.height_rule = rh, WD_ROW_HEIGHT_RULE.AT_LEAST
    for r in t.rows:
        cell_borders(r.cells[4], ("top", "bottom"))
    last = t.rows[n + 1]
    a = last.cells[5].merge(last.cells[6])
    P([("Итого:", True)], align=A.RIGHT, cell=a)
    P([(money(d["schedule_total"]) if sch else "", True)], align=A.RIGHT, cell=last.cells[7])
    for k in range(5):
        cell_borders(last.cells[k], ("left", "bottom", "right"))

    H("4. Обязательства сторон")
    items = ["Покупатель подтверждает получение указанного товара в надлежащем состоянии, согласен с его стоимостью и обязуется вносить платежи в сроки и суммах согласно графику.",
             "Если Покупатель пропустит два платежа подряд, Продавец вправе потребовать погашения всего оставшегося долга сразу.",
             "При неоплате споры разрешаются в суде по месту нахождения Продавца в соответствии с законодательством Кыргызской Республики.",
             "Покупатель обязуется сообщить Продавцу о смене телефона или адреса. Уведомления, направленные по указанному телефону (звонок, SMS, WhatsApp), считаются полученными."]
    for i, x in enumerate(items, 1):
        P(f"4.{i}. {x}", after=1).alignment = A.JUSTIFY
    H("5. Подтверждение и подписи")
    P("С условиями ознакомлен(а), товар получил(а), подписываю добровольно.", after=8)
    t = doc.add_table(rows=3, cols=2); borders(t, False); widths(t, [9, 9])
    P([("Покупатель:", True)], cell=t.cell(0, 0)); P([("Продавец: ", True), (SELLER_NAME, False)], cell=t.cell(0, 1))
    P([("_____________ /", False), V(d["fio"], "______________________"), ("/", False)], cell=t.cell(1, 0), before=10)
    P("_____________ /______________________/", cell=t.cell(1, 1), before=10)
    P("       подпись                   Ф. И. О.", cell=t.cell(2, 0), size=8)
    P("       подпись                   Ф. И. О.", cell=t.cell(2, 1), size=8)
    P("Дата подписания: «____» ______________ 20___ г.", before=8, after=6)
    P("Составлено в двух экземплярах, по одному для каждой стороны.", size=9.5).runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- Streamlit ----------------
def load_sale_data(sale_id):
    from database import supabase
    sale = (supabase.table("sales").select("*").eq("id", sale_id).execute().data or [None])[0]
    if not sale:
        raise ValueError("Договор не найден")
    client = {}
    if sale.get("client_id") is not None:
        client = (supabase.table("clients").select("*").eq("id", sale["client_id"]).execute().data or [{}])[0]
    payments = supabase.table("credit_payments").select("*").eq("sale_id", sale_id).execute().data or []
    return build_obligation_data(sale, client, payments, seller_settings())


def show_obligation(sale_id, key="obl"):
    """Бланк на экране (с кнопкой «🖨 Печать») + «Скачать Word»."""
    import streamlit as st
    import streamlit.components.v1 as components
    try:
        data = load_sale_data(sale_id)
    except Exception as e:
        st.error(f"Бланк не получилось собрать: {e}")
        return
    st.download_button("📝 Скачать Word", data=build_docx(data), file_name=f"Obyazatelstvo_{data['number']}.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                       key=f"{key}_docx_{sale_id}", use_container_width=True)
    st.caption("«🖨 Печать» открывает печать бланка. С телефона удобнее скачать Word и распечатать.")
    rows = schedule_layout(len(data["schedule"]))
    components.html(render_html(data), height=1250 + max(0, rows - 6) * 25, scrolling=True)


def obligation_button(sale_id, key="obl"):
    """Кнопка «📄 Бланк обязательства», под ней — бланк (переключатель)."""
    import streamlit as st
    flag = f"{key}_open_{sale_id}"
    if st.button("📄 Бланк обязательства", key=f"{key}_btn_{sale_id}", use_container_width=True):
        st.session_state[flag] = not st.session_state.get(flag, False)
    if st.session_state.get(flag):
        show_obligation(sale_id, key=key)


DOC_COLUMNS = ("pin", "passport_issued_by", "passport_date")


def client_doc_columns_available():
    """Есть ли в clients колонки из migrations/003_client_documents.sql (проверка один раз за сессию)."""
    import streamlit as st
    if "client_doc_cols" not in st.session_state:
        try:
            from database import supabase
            supabase.table("clients").select(",".join(DOC_COLUMNS)).limit(1).execute()
            st.session_state["client_doc_cols"] = True
        except Exception:
            st.session_state["client_doc_cols"] = False
    return st.session_state["client_doc_cols"]


def doc_fields_input(person=None, prefix="new"):
    """Поля ПИН / кем выдан / дата выдачи в форме клиента — только если миграция выполнена."""
    import streamlit as st
    if not client_doc_columns_available():
        return {}
    person = person or {}
    pin = st.text_input("ПИН (14 цифр, необязательно)", value=str(person.get("pin") or ""), key=f"{prefix}_pin").strip()
    issued = st.text_input("Паспорт: кем выдан (необязательно)", value=str(person.get("passport_issued_by") or ""),
                           key=f"{prefix}_issued").strip()
    cur = _day(person.get("passport_date"))
    ptxt = st.text_input("Дата выдачи паспорта, ДД.ММ.ГГГГ (необязательно)", value=cur.strftime("%d.%m.%Y") if cur else "",
                         key=f"{prefix}_pdate").strip()
    pdate = _day(ptxt) if ptxt else None
    return {"pin": pin or None, "passport_issued_by": issued or None,
            "passport_date": pdate.strftime("%Y-%m-%d") if pdate else None}
